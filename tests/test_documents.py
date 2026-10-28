"""Markdown / HTML / Word conversions."""
import pymupdf
import pytest
from docx import Document
from PIL import Image

from dragshift import convert, documents
from dragshift.util import UserError

MD = """# Trip plan

Some **bold** and *italic* text with `code`.

- first
- second
  - nested

1. one
2. two

| Day | Place |
| --- | ----- |
| Mon | Paris |

> a quote

```
raw  block
```
"""


def pdf_text(p):
    with pymupdf.open(p) as d:
        return "\n".join(pg.get_text() for pg in d)


def docx_paras(p):
    return [(x.style.name, x.text) for x in Document(p).paragraphs]


def test_targets_depend_on_the_file_type():
    assert convert.targets_for("text", ".md") == ["PDF", "DOCX", "HTML"]
    assert convert.targets_for("text", ".HTML") == ["PDF", "DOCX", "MD", "TXT"]
    assert convert.targets_for("text", ".htm") == ["PDF", "DOCX", "MD", "TXT"]
    assert convert.targets_for("text", ".txt") == ["PDF", "DOCX"]          # unchanged
    assert convert.targets_for("docx", ".docx") == ["PDF", "TXT", "MD", "HTML"]


def test_markdown_to_pdf_is_rendered_not_raw_source(tmp_path):
    src = tmp_path / "plan.md"
    src.write_text(MD, encoding="utf-8")
    (out,) = convert.convert(src, "text", "PDF")
    text = pdf_text(out)
    for word in ("Trip plan", "Paris", "nested", "a quote", "raw  block"):
        assert word in text
    assert "# Trip plan" not in text and "**bold**" not in text and "| Day |" not in text


def test_markdown_to_docx_keeps_structure(tmp_path):
    src = tmp_path / "plan.md"
    src.write_text(MD, encoding="utf-8")
    (out,) = convert.convert(src, "text", "DOCX")
    paras = docx_paras(out)
    assert ("Heading 1", "Trip plan") in paras
    assert ("List Bullet", "first") in paras and ("List Bullet 2", "nested") in paras
    assert ("List Number", "one") in paras
    assert ("Quote", "a quote") in paras
    doc = Document(out)
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Day", "Place"]
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Mon", "Paris"]
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "bold" in bold


def test_markdown_to_html_file(tmp_path):
    src = tmp_path / "plan.md"
    src.write_text(MD, encoding="utf-8")
    (out,) = convert.convert(src, "text", "HTML")
    page = out.read_text(encoding="utf-8")
    assert page.startswith("<!doctype html>") and "<h1>Trip plan</h1>" in page and "<table>" in page


HTML = """<html><head><title>T</title><style>p{color:red}</style><script>alert(1)</script></head>
<body><h2>Report</h2><p>Hello <b>world</b> and <a href="http://x.y">a link</a>.</p>
<ul><li>alpha</li><li>beta<ul><li>gamma</li></ul></li></ul>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>"""


def test_html_to_pdf_shows_the_rendered_page_without_tags_or_scripts(tmp_path):
    src = tmp_path / "page.html"
    src.write_text(HTML, encoding="utf-8")
    (out,) = convert.convert(src, "text", "PDF")
    text = pdf_text(out)
    assert "Report" in text and "world" in text and "gamma" in text
    assert "<h2>" not in text and "alert" not in text and "color:red" not in text


def test_html_to_docx_has_no_duplicated_nested_list_text(tmp_path):
    src = tmp_path / "page.html"
    src.write_text(HTML, encoding="utf-8")
    (out,) = convert.convert(src, "text", "DOCX")
    paras = docx_paras(out)
    assert ("Heading 2", "Report") in paras
    texts = [t for _, t in paras]
    assert texts.count("gamma") == 1 and texts.count("beta") == 1
    assert ("List Bullet 2", "gamma") in paras


def test_html_to_md_and_txt(tmp_path):
    src = tmp_path / "page.htm"
    src.write_text(HTML, encoding="utf-8")
    (md,) = convert.convert(src, "text", "MD")
    body = md.read_text(encoding="utf-8")
    assert "## Report" in body and "**world**" in body and "- alpha" in body and "alert" not in body
    (txt,) = convert.convert(src, "text", "TXT")
    plain = txt.read_text(encoding="utf-8")
    assert "Report" in plain and "gamma" in plain and "<" not in plain and "alert" not in plain


def test_relative_images_are_embedded_in_pdf_and_docx(tmp_path):
    Image.new("RGB", (120, 80), (200, 30, 30)).save(tmp_path / "pic.png")
    src = tmp_path / "doc.md"
    src.write_text("# With picture\n\n![alt](pic.png)\n\nafter\n", encoding="utf-8")
    (pdf,) = convert.convert(src, "text", "PDF")
    with pymupdf.open(pdf) as d:
        assert len(d[0].get_images()) == 1
    (dx,) = convert.convert(src, "text", "DOCX")
    assert len(Document(dx).inline_shapes) == 1


def test_remote_and_missing_images_are_skipped_without_failing(tmp_path):
    src = tmp_path / "doc.md"
    src.write_text("![a](http://example.com/x.png)\n\n![b](nope.png)\n\ntext\n", encoding="utf-8")
    convert.convert(src, "text", "DOCX")
    convert.convert(src, "text", "PDF")


def make_docx(path):
    d = Document()
    d.add_heading("Title here", 1)
    d.add_heading("Sub", 2)
    p = d.add_paragraph("Plain then ")
    r = p.add_run("bold")
    r.bold = True
    p.add_run(" and ").italic = False
    p.add_run("slanted").italic = True
    d.add_paragraph("item a", style="List Bullet")
    d.add_paragraph("item b", style="List Bullet")
    d.add_paragraph("step", style="List Number")
    t = d.add_table(rows=2, cols=2)
    for i, row in enumerate([["H1", "H2"], ["x|y", "z"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph("Last paragraph")
    d.save(path)


def test_docx_to_markdown(tmp_path):
    src = tmp_path / "in.docx"
    make_docx(src)
    (out,) = convert.convert(src, "docx", "MD")
    md = out.read_text(encoding="utf-8")
    assert md.startswith("# Title here\n\n## Sub\n\nPlain then **bold** and *slanted*\n\n")
    assert "- item a\n- item b\n\n1. step" in md
    assert "| H1 | H2 |\n| --- | --- |\n| x\\|y | z |" in md
    assert md.endswith("Last paragraph\n")


def test_docx_to_markdown_round_trips_through_markdown_rendering(tmp_path):
    src = tmp_path / "in.docx"
    make_docx(src)
    (md,) = convert.convert(src, "docx", "MD")
    page = documents.md_to_html(md.read_text(encoding="utf-8"))
    assert "<h1>Title here</h1>" in page and "<strong>bold</strong>" in page
    assert page.count("<li>") == 3 and "<table>" in page


def test_docx_to_html_file(tmp_path):
    src = tmp_path / "in.docx"
    make_docx(src)
    (out,) = convert.convert(src, "docx", "HTML")
    page = out.read_text(encoding="utf-8")
    assert page.startswith("<!doctype html>") and "<h1>Title here</h1>" in page and "<b>bold</b>" in page


def test_unsupported_target_and_plain_text_are_unchanged(tmp_path):
    src = tmp_path / "a.md"
    src.write_text("hi", encoding="utf-8")
    with pytest.raises(UserError):
        documents.convert_markup(src, "TXT")
    t = tmp_path / "n.txt"
    t.write_text("# not markdown", encoding="utf-8")
    (pdf,) = convert.convert(t, "text", "PDF")
    assert "# not markdown" in pdf_text(pdf)         # .txt stays literal
